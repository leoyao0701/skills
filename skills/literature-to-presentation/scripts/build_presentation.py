#!/usr/bin/env python3
"""Render a source-grounded JSON deck to HTML, paired notes and optional native PPTX."""
import argparse
from html import escape
import json
from pathlib import Path
import re
import math
import unicodedata

BASE = Path(__file__).resolve().parents[1]
LAYOUTS = {'text', 'parallel', 'comparison', 'timeline', 'matrix', 'argument'}


def check(data):
    errors, warnings = [], []
    report = {'errors': errors, 'warnings': warnings,
              'scope': '结构、字段类型、引用关联与密度提示；不验证引文真实性、因果关系或实际渲染。'}
    if not isinstance(data, dict):
        errors.append('顶层必须是对象。')
        return report

    def text_fields(obj, fields, label):
        for field in fields:
            if field in obj and not isinstance(obj[field], str):
                errors.append(f'{label}的 {field} 必须是字符串。')

    text_fields(data, ['title', 'language'], '文档')
    if not isinstance(data.get('title'), str) or not data['title'].strip():
        errors.append('需要非空 title。')
    sources = data.get('sources', [])
    if not isinstance(sources, list) or any(not isinstance(src, dict) for src in sources):
        errors.append('sources 必须是来源对象列表。')
        sources = []
    known = {}
    for src in sources:
        text_fields(src, ['id', 'citation', 'locator', 'scope', 'quote'], '来源')
        sid = src.get('id')
        if not isinstance(sid, str) or not sid.strip() or sid in known:
            errors.append('来源 ID 为空、类型无效或重复。')
        else:
            known[sid] = src
    slides = data.get('slides', [])
    if not isinstance(slides, list) or not slides or any(not isinstance(s, dict) for s in slides):
        errors.append('slides 必须是非空页面对象列表。')
        return report
    seen = set()
    for i, s in enumerate(slides, 1):
        label = f'第 {i} 页'
        text_fields(s, ['id', 'title', 'layout', 'body', 'relation', 'notes', 'quote_source'], label)
        sid = s.get('id', '')
        if not isinstance(sid, str) or not re.fullmatch(r'[A-Za-z][A-Za-z0-9_-]*', sid) or sid in seen:
            errors.append(f'{label} ID 无效或重复。')
        seen.add(str(sid))
        layout = s.get('layout', 'text')
        if not isinstance(layout, str) or layout not in LAYOUTS:
            errors.append(f'{label}版式无效。')
            layout = 'text'
        if not isinstance(s.get('title'), str) or not s['title'].strip():
            errors.append(f'{label}缺少标题。')
        refs = s.get('sources', [])
        if not isinstance(refs, list) or any(not isinstance(r, str) for r in refs):
            errors.append(f'{label} sources 必须是来源 ID 字符串列表。')
            refs = []
        if any(r not in known for r in refs):
            errors.append(f'{label}引用了不存在的来源。')
        quote = s.get('quote_source')
        if quote:
            if not isinstance(quote, str) or quote not in refs or quote not in known:
                errors.append(f'{label} quote_source 必须属于本页 sources。')
            elif not isinstance(known[quote].get('quote'), str) or not known[quote]['quote'].strip():
                errors.append(f'{label}展示的来源必须提供非空 quote。')
        if layout != 'text' and not s.get('relation'):
            errors.append(f'{label}需声明 relation，说明组件表达的关系。')
        if not s.get('notes'):
            warnings.append(f'{label}缺少详细讲稿。')
        items = s.get('items', [])
        if not isinstance(items, list) or any(not isinstance(item, dict) for item in items):
            errors.append(f'{label} items 必须是对象列表。')
            items = []
        for item in items:
            text_fields(item, ['title', 'text'], label + '组件')
        if layout == 'matrix':
            m = s.get('matrix', {})
            if not isinstance(m, dict):
                errors.append(f'{label} matrix 必须是对象。')
                continue
            columns, rows = m.get('columns', []), m.get('rows', [])
            if (not isinstance(columns, list) or len(columns) < 2
                    or any(not isinstance(c, str) for c in columns)
                    or not isinstance(rows, list) or not rows
                    or any(not isinstance(row, list) or len(row) != len(columns)
                           or any(not isinstance(c, str) for c in row) for row in rows)):
                errors.append(f'{label}矩阵至少两列；列名和单元格须为字符串，各行须为等长列表。')
            elif len(columns) > 4 or len(rows) > 5:
                warnings.append(f'{label}矩阵较大，请拆页并检查可读性。')
            if items:
                errors.append(f'{label}矩阵版式不能同时使用 items。')
        elif layout != 'text' and not items:
            errors.append(f'{label}该版式需要 items。')
        if len(items) > 4:
            warnings.append(f'{label}组件超过四项，请检查密度。')
        texts = [s.get('body', ''), *(x.get('text', '') for x in items)]
        if sum(len(t) for t in texts if isinstance(t, str)) > 480:
            warnings.append(f'{label}正文较长；这是密度提示，不是溢出判定。')
    return report


def paragraphs(text):
    return ''.join('<p>' + escape(p).replace('\n', '<br>') + '</p>' for p in text.split('\n\n') if p)


def body_html(s, sources):
    result = paragraphs(s.get('body', ''))
    if s.get('quote_source'):
        src = sources[s['quote_source']]
        result += '<blockquote>' + escape(src.get('quote', '')) + '</blockquote>'
    if s.get('layout') == 'matrix':
        m = s['matrix']
        result += '<table><thead><tr>' + ''.join('<th>' + escape(c) + '</th>' for c in m['columns']) + '</tr></thead><tbody>'
        result += ''.join('<tr>' + ''.join('<td>' + escape(c) + '</td>' for c in row) + '</tr>' for row in m['rows']) + '</tbody></table>'
    elif s.get('items'):
        layout = s.get('layout', 'text')
        result += f'<div class="components {layout}">'
        for item in s['items']:
            result += '<section class="component"><h3>' + escape(item.get('title', '')) + '</h3>' + paragraphs(item.get('text', '')) + '</section>'
        result += '</div>'
    if s.get('relation'):
        result += '<p class="relation">' + escape(s['relation']) + '</p>'
    return result


def source_text(source):
    return ' · '.join(str(source.get(k, '')) for k in ['citation', 'locator', 'scope'] if source.get(k))


def notes_text(s, sources):
    originals = '\n\n'.join(source.get('quote', '') + '\n' + source_text(source)
                            for source in (sources[r] for r in s.get('sources', [])))
    return s.get('notes', '') + ('\n\n原文与来源\n' + originals if originals else '')


def write_html(data, out):
    sources = {s['id']: s for s in data.get('sources', [])}
    css = (BASE / 'assets/components.css').read_text(encoding='utf-8')
    js = (BASE / 'assets/navigation.js').read_text(encoding='utf-8')
    slides, pairs, md = [], [], ['# ' + data['title'], '']
    for i, s in enumerate(data['slides'], 1):
        refs = [sources[r] for r in s.get('sources', [])]
        slides.append('<article class="slide" id="' + escape(s['id']) + '"><header><span class="number">' + str(i).zfill(2) + '</span><h1>' + escape(s['title']) + '</h1></header><div class="content">' + body_html(s, sources) + '</div><footer>' + escape(data['title']) + '</footer></article>')
        originals = ''.join('<figure><blockquote>' + escape(r.get('quote', '')) + '</blockquote><figcaption>' + escape(source_text(r)) + '</figcaption></figure>' for r in refs)
        pairs.append('<article class="spread"><h2>' + str(i) + ' · ' + escape(s['title']) + '</h2><div class="pair"><aside>' + (originals or '<p>本页未提供原文。</p>') + '</aside><section>' + paragraphs(s.get('notes', '')) + '</section></div></article>')
        md += ['## ' + str(i) + ' · ' + s['title'], '', notes_text(s, sources), '']
    head = '<!doctype html><html lang="' + escape(data.get('language', 'zh-CN'), quote=True) + '"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>' + escape(data['title']) + '</title><style>' + css + '</style></head><body>'
    nav = '<nav><button data-action="prev" aria-label="上一页">←</button><span data-role="position"></span><button data-action="next" aria-label="下一页">→</button><button data-action="show-all">连续阅读</button><button data-action="fullscreen">全屏</button><a href="notes.html">对照讲稿</a></nav>'
    (out / 'index.html').write_text(head + '<main>' + ''.join(slides) + '</main>' + nav + '<script>' + js + '</script></body></html>', encoding='utf-8')
    (out / 'notes.html').write_text(head + '<main class="notes">' + ''.join(pairs) + '</main></body></html>', encoding='utf-8')
    (out / 'notes.md').write_text('\n'.join(md), encoding='utf-8')
    (out / 'sources.json').write_text(json.dumps(data.get('sources', []), ensure_ascii=False, indent=2), encoding='utf-8')


def write_pptx(data, dest, font):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
    except ImportError as exc:
        raise RuntimeError('PPTX 需要 python-pptx；请安装 scripts/requirements-pptx.txt。') from exc
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333333), Inches(7.5)
    sources = {s['id']: s for s in data.get('sources', [])}
    ink, pale = '233B35', 'EEF1E9'

    def required_height(text, width, size, padding=.08):
        # Conservative estimate; actual fonts and word wrapping still need visual review.
        capacity = max(1, (width - .2) * 72 / size)
        lines = text.split('\n')
        count = sum(max(1, math.ceil(sum(1 if unicodedata.east_asian_width(c) in 'WF' else .6
                                       for c in line) / capacity)) for line in lines)
        return (count * size * 1.2 + max(0, len(lines) - 1) * 6) / 72 + padding * 2

    def box(slide, text, x, y, w, h, size=21, bold=False, fill=None, padding=.08):
        if required_height(text, w, size, padding) > h + .01:
            raise RuntimeError(f'第 {len(p.slides)} 页文本估算超出可用空间，请拆页或调整布局：{text[:36]}')
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        from pptx.enum.text import MSO_AUTO_SIZE
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = Inches(.1)
        tf.margin_top = tf.margin_bottom = Inches(padding)
        if fill:
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        for j, line in enumerate(text.split('\n')):
            paragraph = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            paragraph.text = line
            paragraph.font.name, paragraph.font.size = font, Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(ink)
            paragraph.line_spacing = 1.2
            paragraph.space_after = Pt(6 if j < len(text.split('\n')) - 1 else 0)
        return shape

    for i, s in enumerate(data['slides'], 1):
        slide = p.slides.add_slide(p.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string('FBFAF6')
        box(slide, f'{i:02d}  {s["title"]}', .6, .35, 12.1, .85, 28, True)
        y = 1.35
        blocks = [(s.get('body', ''), None)]
        if s.get('quote_source'):
            blocks.append((sources[s['quote_source']]['quote'], pale))
        for text, fill in blocks:
            if text:
                h = max(.75, required_height(text, 12, 20))
                if y + h > 6.35:
                    raise RuntimeError(f'第 {i} 页正文或引文过长，请拆页；未写入新的 PPTX。')
                box(slide, text, .65, y, 12, h, 20, fill=fill)
                y += h + .15
        available = 6.35 - y
        layout = s.get('layout', 'text')
        if layout == 'matrix':
            matrix = s['matrix']; rows = [matrix['columns'], *matrix['rows']]
            heights = [max(required_height(t, 12 / len(row), 18) for t in row) for row in rows]
            if sum(heights) > available:
                raise RuntimeError(f'第 {i} 页表格估算过密，请拆页。')
            table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(.65), Inches(y), Inches(12), Inches(available)).table
            extra = (available - sum(heights)) / len(rows)
            for r, row in enumerate(rows):
                table.rows[r].height = Inches(heights[r] + extra)
                for c, text in enumerate(row):
                    cell = table.cell(r, c); cell.text = text
                    cell.margin_left = cell.margin_right = Inches(.1)
                    cell.margin_top = cell.margin_bottom = Inches(.08)
                    cell.text_frame.word_wrap = True
                    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(pale if r == 0 else 'FBFAF6')
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.name = font; paragraph.font.size = Pt(18)
                        paragraph.line_spacing = 1.2; paragraph.space_after = Pt(0)
                        paragraph.font.bold = r == 0; paragraph.font.color.rgb = RGBColor.from_string(ink)
        elif s.get('items'):
            items = s['items']; horizontal = layout in {'parallel', 'comparison'}
            width = (12 - .2 * (len(items) - 1)) / len(items) if horizontal else 12
            if width <= .2:
                raise RuntimeError(f'第 {i} 页组件过多，请拆页。')
            title_heights = [max(.48, required_height(item.get('title', ''), width, 19)) for item in items]
            heights = [th + required_height(item.get('text', ''), width, 19) for th, item in zip(title_heights, items)]
            need = max(heights) if horizontal else sum(heights) + .15 * (len(items) - 1)
            if need > available:
                raise RuntimeError(f'第 {i} 页组件估算过密，请拆页或调整布局。')
            top = y
            for j, item in enumerate(items):
                th = max(title_heights) if horizontal else title_heights[j]
                if horizontal:
                    x, h = .65 + j * (width + .2), available
                else:
                    x, h = .65, heights[j] + (available - need) / len(items)
                box(slide, item.get('title', ''), x, top, width, th, 19, True, pale)
                box(slide, item.get('text', ''), x, top + th, width, h - th, 19, fill=pale)
                if not horizontal:
                    top += h + .15
        box(slide, s.get('relation', ''), .65, 6.55, 12, .43, 13)
        box(slide, data['title'], .65, 7.08, 12, .22, 9, padding=0)
        slide.notes_slide.notes_text_frame.text = notes_text(s, sources)
    p.save(dest)
    # Reopen the produced file, not the in-memory draft.
    q = Presentation(dest)
    findings = []
    for i, slide in enumerate(q.slides, 1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > q.slide_width + 10 or shape.top + shape.height > q.slide_height + 10:
                findings.append(f'第 {i} 页有超出画布的形状。')
        expected = notes_text(data['slides'][i-1], sources)
        if slide.notes_slide.notes_text_frame.text != expected:
            findings.append(f'第 {i} 页讲稿备注不同步。')
    return {'slides': len(q.slides), 'errors': findings,
            'scope': '导出前估算文本空间；最终 PPTX 可重新读取、形状边界、备注一致性。估算不代替字体与溢出的视觉检查，也不验证原文真实性。'}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('input', type=Path)
    ap.add_argument('--out', type=Path, required=True)
    ap.add_argument('--pptx', action='store_true')
    ap.add_argument('--font', default='Noto Sans CJK SC')
    args = ap.parse_args()
    data = json.loads(args.input.read_text(encoding='utf-8'))
    report = check(data)
    if report['errors']:
        raise SystemExit('\n'.join(report['errors']))
    args.out.mkdir(parents=True, exist_ok=True)
    write_html(data, args.out)
    report['formats'] = ['html', 'notes.html', 'notes.md', 'sources.json']
    if args.pptx:
        try:
            report['pptx'] = write_pptx(data, args.out / 'presentation.pptx', args.font)
        except RuntimeError as exc:
            report['pptx'] = {'errors': [str(exc)], 'scope': 'PPTX 导出未完成；输出目录中如有旧 PPTX，不代表本轮结果。'}
        if not report['pptx']['errors']:
            report['formats'].append('pptx')
    (args.out / 'checks.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False))
    if report.get('pptx', {}).get('errors'):
        raise SystemExit(1)


if __name__ == '__main__':
    main()
