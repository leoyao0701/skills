"""Run from any directory: python3 -m unittest discover -s <skill>/tests -v."""
import copy
import importlib.util
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE / 'scripts'))
import build_presentation as build
import extract_pdf


def demo():
    return json.loads((BASE / 'examples/demo.json').read_text(encoding='utf-8'))


class ContentTests(unittest.TestCase):
    def test_example_and_source_sync(self):
        data = demo()
        self.assertEqual(build.check(data)['errors'], [])
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)
            build.write_html(data, out)
            self.assertEqual(json.loads((out / 'sources.json').read_text()), data['sources'])
            for slide in data['slides']:
                self.assertIn(slide['title'], (out / 'index.html').read_text())
                self.assertIn(slide['notes'], (out / 'notes.md').read_text())
            self.assertIn('会谈记录、交换文件', (out / 'notes.html').read_text())

    def test_bad_container_and_text_types_rejected(self):
        inputs = [[], None, {'title': 'x', 'slides': 'x'}, {'title': 'x', 'slides': [3]}]
        for key, value in [('sources', [3]), ('sources', [{'id': []}]), ('title', 3)]:
            data = demo(); data[key] = value; inputs.append(data)
        for key, value in [('sources', [{}]), ('body', 3), ('layout', []), ('items', [None])]:
            data = demo(); data['slides'][0][key] = value; inputs.append(data)
        for data in inputs:
            with self.subTest(data=data):
                self.assertTrue(build.check(data)['errors'])

    def test_invalid_matrix_rejected(self):
        for rows in [[[2020, 12, 3]], ['xyz'], [['short']], None]:
            data = demo(); data['slides'][4]['matrix']['rows'] = rows
            self.assertTrue(build.check(data)['errors'])

    def test_source_relationships_rejected(self):
        for action in ['unknown', 'unlisted', 'empty_quote', 'duplicate']:
            data = demo()
            if action == 'unknown': data['slides'][0]['sources'] = ['absent']
            elif action == 'unlisted': data['slides'][0]['sources'] = []
            elif action == 'empty_quote': data['sources'][0]['quote'] = ''
            else: data['sources'].append(copy.deepcopy(data['sources'][0]))
            self.assertTrue(build.check(data)['errors'], action)

    def test_html_treats_text_as_text(self):
        data = demo(); payload = '<script>window.bad=true</script>'
        data['slides'][0]['body'] = payload
        data['slides'][0]['notes'] = payload
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp); build.write_html(data, out)
            for name in ['index.html', 'notes.html']:
                text = (out / name).read_text()
                self.assertNotIn(payload, text)
                self.assertIn('&lt;script&gt;window.bad=true&lt;/script&gt;', text)

    def test_pdf_selection(self):
        self.assertEqual(extract_pdf.selected_pages('1,3-5', 5), {1, 3, 4, 5})
        for spec in ['0', '6', '4-2', '1-2-3']:
            with self.assertRaises(ValueError): extract_pdf.selected_pages(spec, 5)


@unittest.skipUnless(importlib.util.find_spec('pptx'), 'optional python-pptx is not installed')
class PowerPointTests(unittest.TestCase):
    def test_editable_table_and_exact_notes(self):
        from pptx import Presentation
        data = demo()
        with tempfile.TemporaryDirectory() as temp:
            dest = Path(temp) / 'demo.pptx'
            self.assertEqual(build.write_pptx(data, dest, 'Arial')['errors'], [])
            slides = Presentation(dest).slides
            self.assertEqual(len(slides), 6)
            table = next(s.table for s in slides[4].shapes if s.has_table)
            self.assertEqual(table.cell(1, 2).text, '会谈记录、交换文件')
            sources = {s['id']: s for s in data['sources']}
            for source, slide in zip(data['slides'], slides):
                self.assertEqual(slide.notes_slide.notes_text_frame.text, build.notes_text(source, sources))

    def test_long_body_does_not_silently_collide(self):
        from pptx import Presentation
        data = demo(); data['slides'] = [data['slides'][0]]
        data['slides'][0]['body'] = '\n'.join('短正文' for _ in range(8))
        with tempfile.TemporaryDirectory() as temp:
            dest = Path(temp) / 'long.pptx'
            build.write_pptx(data, dest, 'Arial')
            shapes = Presentation(dest).slides[0].shapes
            body = next(s for s in shapes if s.has_text_frame and s.text.startswith('短正文'))
            quote = next(s for s in shapes if s.has_text_frame and s.text.startswith('协商可能'))
            self.assertLessEqual(body.top + body.height, quote.top)
            data['slides'][0]['body'] *= 10
            with self.assertRaisesRegex(RuntimeError, '拆页'):
                build.write_pptx(data, Path(temp) / 'excess.pptx', 'Arial')
            self.assertFalse((Path(temp) / 'excess.pptx').exists())


@unittest.skipUnless(importlib.util.find_spec('pypdf'), 'optional pypdf is not installed')
class PDFTests(unittest.TestCase):
    def test_empty_page_is_not_verified_and_ocr_missing_is_explicit(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as temp:
            pdf, out = Path(temp) / 'sample.pdf', Path(temp) / 'sample.json'
            writer = PdfWriter(); writer.add_blank_page(width=300, height=400)
            with pdf.open('wb') as stream: writer.write(stream)
            args = ['extract_pdf.py', str(pdf), '--out', str(out)]
            with patch.object(sys, 'argv', args): extract_pdf.main()
            page = json.loads(out.read_text())['pages'][0]
            self.assertEqual(page['pdf_page'], 1)
            self.assertIsNone(page['printed_page'])
            self.assertEqual(page['review_status'], 'not_verified')
            self.assertIn('warning', page)
            with patch.object(sys, 'argv', [*args, '--ocr-pages', '1']), patch('extract_pdf.shutil.which', return_value=None):
                with self.assertRaisesRegex(SystemExit, 'Tesseract'): extract_pdf.main()


@unittest.skipUnless(os.environ.get('LTP_BROWSER_EXECUTABLE') and importlib.util.find_spec('playwright'),
                     'set LTP_BROWSER_EXECUTABLE and install Playwright for isolated browser test')
class BrowserTests(unittest.TestCase):
    def test_reserved_ids_and_notes_layout(self):
        from playwright.sync_api import sync_playwright
        data = demo()
        for slide, sid in zip(data['slides'], ['position', 'prev', 'next', 'all', 'fullscreen', 'notes']):
            slide['id'] = sid
        with tempfile.TemporaryDirectory() as temp, sync_playwright() as pw:
            out = Path(temp); build.write_html(data, out)
            browser = pw.chromium.launch(executable_path=os.environ['LTP_BROWSER_EXECUTABLE'])
            page = browser.new_page(viewport={'width': 1600, 'height': 1100})
            page.goto((out / 'index.html').as_uri())
            self.assertEqual(page.locator('#position h1').inner_text(), data['slides'][0]['title'])
            page.locator('nav [data-action="next"]').click()
            self.assertTrue(page.locator('#prev').is_visible())
            page.goto((out / 'notes.html').as_uri())
            pair = page.locator('.pair').first
            left, right = pair.locator('aside').bounding_box(), pair.locator('section').bounding_box()
            self.assertLess(left['x'] + left['width'], right['x'])
            self.assertIn('会谈记录、交换文件', page.locator('body').inner_text())
            browser.close()


if __name__ == '__main__':
    unittest.main()
